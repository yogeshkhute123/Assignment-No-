"""Write the audit as Excel, HTML and JSON."""
import html
import json
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CATEGORIES, SEVERITY_ORDER

SEV_FILL = {"High": "F8CBAD", "Medium": "FFE699", "Low": "DDEBF7"}


def group(issues):
    """One row per (category, code): the issue, its recommendation, and every affected URL."""
    g = {}
    for i in issues:
        k = (i.category, i.code)
        if k not in g:
            g[k] = {"category": i.category, "severity": i.severity, "title": i.title,
                    "recommendation": i.recommendation, "urls": [], "details": []}
        if i.url and i.url not in g[k]["urls"]:
            g[k]["urls"].append(i.url)
        if i.detail:
            g[k]["details"].append(f"{i.url}: {i.detail}" if i.url else i.detail)
    return sorted(g.values(), key=lambda r: (CATEGORIES.index(r["category"]), SEVERITY_ORDER[r["severity"]], -len(r["urls"])))


def _page_rows(site):
    rows = []
    for p in site.pages.values():
        s = p.soup if p.is_html else None
        title = s.find("title").get_text(strip=True) if s and s.find("title") else ""
        h1 = s.find("h1").get_text(" ", strip=True) if s and s.find("h1") else ""
        rows.append([p.url, p.status or p.error, p.depth, round(p.elapsed, 2), p.size, title, h1,
                     len(site.inlinks.get(p.url, [])), len(p.internal_links), len(p.external_links)])
    return rows


def _style_sheet(ws, widths):
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_xlsx(path, site, issues, scores):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["SEO Audit", site.start_url])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append(["Pages crawled", len(site.pages)] + (["(crawl limit reached - raise --max-pages)"] if site.truncated else []))
    ws.append([])
    ws.append(["Category", "Score /100", "High", "Medium", "Low", "Total issues"])
    for cat in CATEGORIES:
        mine = [i for i in issues if i.category == cat]
        ws.append([cat, scores[cat]] + [sum(1 for i in mine if i.severity == s) for s in ("High", "Medium", "Low")] + [len(mine)])
    ws.append(["Overall", scores["Overall"]])
    for c in ws[5]:
        c.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 45

    def issues_sheet(name, rows):
        sh = wb.create_sheet(name)
        sh.append(["Sr No", "Category", "Severity", "Issue", "Recommendation", "Affected URLs", "Count", "Details"])
        for n, r in enumerate(rows, 1):
            sh.append([n, r["category"], r["severity"], r["title"], r["recommendation"],
                       "\n".join(r["urls"][:200]) + (f"\n… +{len(r['urls'])-200} more" if len(r["urls"]) > 200 else ""),
                       len(r["urls"]) or 1, "\n".join(r["details"][:50])])
            sh.cell(n + 1, 3).fill = PatternFill("solid", fgColor=SEV_FILL[r["severity"]])
        _style_sheet(sh, [7, 18, 10, 38, 55, 60, 8, 70])

    grouped = group(issues)
    issues_sheet("All Issues", grouped)
    for cat in CATEGORIES:
        issues_sheet(cat.split(" (")[0], [r for r in grouped if r["category"] == cat])

    sh = wb.create_sheet("Issues per URL")
    sh.append(["Category", "Severity", "Issue", "URL", "Detail", "Recommendation"])
    for i in issues:
        sh.append([i.category, i.severity, i.title, i.url, i.detail, i.recommendation])
    _style_sheet(sh, [18, 10, 38, 60, 70, 55])

    sh = wb.create_sheet("Pages")
    sh.append(["URL", "Status", "Depth", "Response (s)", "Bytes", "Title", "H1", "Inlinks", "Internal links", "External links"])
    for r in _page_rows(site):
        sh.append(r)
    _style_sheet(sh, [60, 9, 7, 12, 10, 45, 45, 9, 14, 14])
    wb.save(path)


def write_html(path, site, issues, scores):
    e = html.escape
    color = lambda s: "#2e7d32" if s >= 80 else "#ef6c00" if s >= 50 else "#c62828"
    cards = "".join(f'<div class="card"><div class="n" style="color:{color(scores[c])}">{scores[c]}</div><div>{e(c)}</div></div>'
                    for c in CATEGORIES + ["Overall"])
    body = []
    for cat in CATEGORIES:
        rows = [r for r in group(issues) if r["category"] == cat]
        body.append(f"<h2>{e(cat)} <small>({len(rows)} issue types)</small></h2>")
        if not rows:
            body.append("<p>No issues found.</p>")
        for r in rows:
            urls = "".join(f"<li><a href='{e(u)}'>{e(u)}</a></li>" for u in r["urls"][:25])
            more = f"<li>… +{len(r['urls'])-25} more (see Excel)</li>" if len(r["urls"]) > 25 else ""
            det = "".join(f"<li>{e(d[:300])}</li>" for d in r["details"][:5])
            body.append(f"<details><summary><span class='sev {r['severity']}'>{r['severity']}</span> {e(r['title'])} "
                        f"<b>({len(r['urls']) or 1})</b></summary><p><b>Fix:</b> {e(r['recommendation'])}</p>"
                        f"<ul>{urls}{more}</ul>{'<p><b>Details</b></p><ul>'+det+'</ul>' if det else ''}</details>")
    doc = f"""<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>SEO Audit - {e(site.start_url)}</title>
<style>body{{font:15px system-ui,sans-serif;max-width:1000px;margin:2em auto;padding:0 16px;color:#222}}
.cards{{display:flex;gap:12px;flex-wrap:wrap}}.card{{border:1px solid #ddd;border-radius:8px;padding:12px 18px;text-align:center}}
.n{{font-size:32px;font-weight:700}}details{{border:1px solid #e0e0e0;border-radius:6px;padding:8px 12px;margin:6px 0}}
summary{{cursor:pointer}}.sev{{padding:1px 8px;border-radius:10px;font-size:12px;margin-right:6px}}
.High{{background:#f8cbad}}.Medium{{background:#ffe699}}.Low{{background:#ddebf7}}small{{color:#777;font-weight:400}}
ul{{word-break:break-all}}</style>
<h1>SEO Audit</h1><p>{e(site.start_url)} · {len(site.pages)} pages crawled · {datetime.now():%Y-%m-%d %H:%M}
{'· <b>crawl limit reached</b>' if site.truncated else ''}</p><div class="cards">{cards}</div>{''.join(body)}"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(doc)


def write_json(path, site, issues, scores):
    data = {"site": site.start_url, "pages_crawled": len(site.pages), "scores": scores,
            "issues": group(issues)}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2)
